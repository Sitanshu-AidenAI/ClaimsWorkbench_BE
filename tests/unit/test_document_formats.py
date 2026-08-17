"""The format readers: PDF, Word, Excel, and saved email.

Every fixture is built in memory from a library this project already depends on, so
there are no binary test assets to keep in step with the code that reads them.

Two things these tests care about beyond "did any text come out":

* **Pages.** A citation is only clickable if the reader said which page a value came
  from, and `char_start`/`char_end` have to address the joined text exactly.
* **Table shape.** A claim form's two-column table has to come out as
  `Label: value`, because that is the one form `app/domain/heuristics.py` can read —
  so several tests assert on the *deterministic reader's* output rather than on the
  extracted text, which is the property that actually matters.
"""

from __future__ import annotations

import datetime as dt
import io
import struct
import zipfile
from email.message import EmailMessage

import pytest

from app.domain.enums import DocumentExtractionStatus
from app.domain.heuristics import extract_from_text
from app.services.documents.formats.tables import render_table
from app.services.documents.text import extract_text

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MSG_TYPE = "application/vnd.ms-outlook"


# --- fixtures -----------------------------------------------------------------


def build_pdf(pages: list[str], *, blank_pages: int = 0) -> bytes:
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    document = canvas.Canvas(buffer)
    for body in pages:
        offset = 800
        for line in body.splitlines():
            document.drawString(60, offset, line)
            offset -= 16
        document.showPage()
    for _ in range(blank_pages):
        document.rect(100, 100, 200, 200, fill=1)
        document.showPage()
    document.save()
    return buffer.getvalue()


def build_docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Description of loss", level=1)
    document.add_paragraph("Water ingress through a failed riser joint.")

    form = document.add_table(rows=3, cols=2)
    for index, (label, value) in enumerate(
        [
            ("Policy number", "CP-2026-4471"),
            ("Insured name", "Harborview Logistics Ltd"),
            ("Date of loss", "08 March 2026"),
        ]
    ):
        form.rows[index].cells[0].text = label
        form.rows[index].cells[1].text = value

    document.add_heading("Schedule", level=1)
    matrix = document.add_table(rows=3, cols=3)
    for row_index, row in enumerate(
        [["Item", "Qty", "Cost"], ["Roof panels", "12", "14190.00"], ["Stock", "1", "38600.00"]]
    ):
        for column_index, value in enumerate(row):
            matrix.rows[row_index].cells[column_index].text = value

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_xlsx() -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = "Policy number"
    summary["B1"] = "CP-2026-4471"
    summary["A2"] = "Date of loss"
    summary["B2"] = dt.datetime(2026, 3, 8)
    # Rows 3-5 blank: three consecutive blanks is past the two-row tolerance, so
    # this is a second table rather than more of the first.
    summary["A6"] = "Adjuster"
    summary["B6"] = "M. Odele"

    schedule = workbook.create_sheet("Schedule of values")
    for row_index, row in enumerate(
        [["Item", "Qty", "Cost"], ["Roof panels", 12, 14190.0], ["Stock", 1, 38600.0]], start=1
    ):
        for column_index, value in enumerate(row, start=1):
            schedule.cell(row=row_index, column=column_index, value=value)
    # Column E is populated far to the right, so D is an empty column inside the
    # used range and must be stripped rather than rendered as a blank column.
    schedule["E1"] = "Notes"
    schedule["E2"] = "Replaced in full"
    schedule["C5"] = "=SUM(C2:C3)"

    hidden = workbook.create_sheet("Internal notes")
    hidden["A1"] = "Do not disclose to the insured"
    hidden.sheet_state = "hidden"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_eml(*, html: bool = True, attachment: bytes | None = None) -> bytes:
    message = EmailMessage()
    message["From"] = "Jane Brooks <jane.brooks@brokerage.example>"
    message["To"] = "claims@carrier.example"
    message["Subject"] = "FNOL - burst pipe at Unit 7"
    message["Date"] = "Tue, 10 Mar 2026 09:14:00 +0000"

    body = (
        "<html><body><p>Policy number: CP-2026-4471</p>"
        "<p>Date of loss: 08 March 2026</p></body></html>"
    )
    if html:
        message.set_content(body, subtype="html")
    else:
        message.set_content("Policy number: CP-2026-4471\nDate of loss: 08 March 2026")

    if attachment is not None:
        message.add_attachment(
            attachment, maintype="application", subtype="pdf", filename="estimate.pdf"
        )
    return message.as_bytes()


# --- a minimal OLE compound file, so the `.msg` reader is tested for real -------

_SECTOR = 512
_MINI_CUTOFF = 4096
_FREE = 0xFFFFFFFF
_ENDOFCHAIN = 0xFFFFFFFE
_FATSECT = 0xFFFFFFFD
_NOSTREAM = 0xFFFFFFFF


def _ole_dir_entry(
    name: str,
    *,
    kind: int,
    right: int = _NOSTREAM,
    child: int = _NOSTREAM,
    start: int = _ENDOFCHAIN,
    size: int = 0,
) -> bytes:
    encoded = name.encode("utf-16-le") + b"\x00\x00"
    entry = encoded.ljust(64, b"\x00")
    entry += struct.pack("<H", len(encoded))
    entry += bytes([kind, 1])  # object type, black
    entry += struct.pack("<III", _NOSTREAM, right, child)
    entry += b"\x00" * 16  # CLSID
    entry += struct.pack("<I", 0)  # state bits
    entry += b"\x00" * 16  # creation and modified time
    entry += struct.pack("<I", start)
    entry += struct.pack("<Q", size)
    return entry


def build_msg(properties: dict[str, str]) -> bytes:
    """An Outlook message file holding the given MAPI string properties.

    `olefile` cannot write, so the container is assembled here. It is a v3 compound
    file with one FAT sector and no mini FAT — every stream is padded past the
    4096-byte mini-stream cutoff, which `olefile` applies whatever the header claims.
    NUL padding is safe because Outlook pads its own property streams and the reader
    strips it.
    """
    entries = [
        (f"__substg1.0_{tag}001F", value.encode("utf-16-le").ljust(_MINI_CUTOFF, b"\x00"))
        for tag, value in properties.items()
    ]

    directory_sectors = max(1, -(-((len(entries) + 1) * 128) // _SECTOR))

    chain: list[int] = []
    payload = b""
    starts: list[int] = []
    sector = 1 + directory_sectors
    for _, data in entries:
        count = -(-len(data) // _SECTOR)
        starts.append(sector)
        chain.extend(
            sector + step + 1 if step < count - 1 else _ENDOFCHAIN for step in range(count)
        )
        payload += data
        sector += count

    directory = [
        _ole_dir_entry("Root Entry", kind=5, child=1 if entries else _NOSTREAM),
        *(
            _ole_dir_entry(
                name,
                kind=2,
                right=index + 2 if index + 1 < len(entries) else _NOSTREAM,
                start=starts[index],
                size=len(data),
            )
            for index, (name, data) in enumerate(entries)
        ),
    ]
    directory_bytes = b"".join(directory).ljust(directory_sectors * _SECTOR, b"\x00")

    table = [_FATSECT]
    table.extend(
        2 + step if step < directory_sectors - 1 else _ENDOFCHAIN
        for step in range(directory_sectors)
    )
    table.extend(chain)
    table += [_FREE] * (_SECTOR // 4 - len(table))
    fat_bytes = b"".join(struct.pack("<I", value) for value in table)

    header = bytearray(_SECTOR)
    header[0:8] = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    struct.pack_into("<H", header, 24, 0x3E)
    struct.pack_into("<H", header, 26, 3)
    struct.pack_into("<H", header, 28, 0xFFFE)
    struct.pack_into("<H", header, 30, 9)
    struct.pack_into("<H", header, 32, 6)
    struct.pack_into("<I", header, 44, 1)
    struct.pack_into("<I", header, 48, 1)
    struct.pack_into("<I", header, 56, _MINI_CUTOFF)
    struct.pack_into("<I", header, 60, _ENDOFCHAIN)
    struct.pack_into("<I", header, 68, _ENDOFCHAIN)
    struct.pack_into("<I", header, 76, 0)
    for slot in range(1, 109):
        struct.pack_into("<I", header, 76 + slot * 4, _FREE)

    return bytes(header) + fat_bytes + directory_bytes + payload


# --- tables -------------------------------------------------------------------


class TestTableRendering:
    def test_a_two_column_form_becomes_labelled_lines(self) -> None:
        rendered = render_table([["Policy number", "CP-2026-4471"], ["Insured", "Harborview Ltd"]])
        assert rendered == "Policy number: CP-2026-4471\nInsured: Harborview Ltd"

    def test_a_matrix_keeps_its_columns(self) -> None:
        rendered = render_table(
            [["Item", "Qty", "Cost"], ["Roof panels", "12", "14190.00"]],
        )
        assert rendered.splitlines()[0] == "| Item | Qty | Cost |"
        assert rendered.splitlines()[1] == "| --- | --- | --- |"

    def test_a_blank_header_falls_back_to_labelled_lines(self) -> None:
        # Three columns, but the header row is mostly empty — so it is not a header,
        # and drawing a grid around it would invent one.
        rendered = render_table([["", "", ""], ["Estimated loss", "GBP", "128000"]])
        assert not rendered.startswith("|")
        assert rendered == "Estimated loss: GBP 128000"

    def test_empty_rows_and_columns_are_dropped(self) -> None:
        rendered = render_table(
            [
                ["Item", "", "Qty", "Cost"],
                ["", "", "", ""],
                ["Roof panels", "", "12", "14190"],
                ["", "", "", ""],
            ]
        )
        assert (
            rendered == "| Item | Qty | Cost |\n| --- | --- | --- |\n| Roof panels | 12 | 14190 |"
        )

    def test_two_columns_always_read_as_a_form_even_with_a_real_header(self) -> None:
        # "Item | Cost" is a genuine header, but two columns cannot be told apart
        # from a label/value form, and getting it wrong in this direction is
        # harmless: "Roof panels: 14190" is still true and still readable. Getting
        # it wrong the other way hides a policy number from the regex reader.
        rendered = render_table([["Item", "Cost"], ["Roof panels", "14190"]])
        assert rendered == "Item: Cost\nRoof panels: 14190"

    def test_a_pipe_inside_a_cell_cannot_close_the_column(self) -> None:
        rendered = render_table([["Item", "Note", "Cost"], ["Roof", "a|b", "14190"]])
        assert "a/b" in rendered
        assert rendered.splitlines()[2].count("|") == 4

    def test_nothing_in_means_nothing_out(self) -> None:
        assert render_table([]) == ""
        assert render_table([["", ""], ["", ""]]) == ""


# --- PDF ----------------------------------------------------------------------


class TestPdf:
    def test_each_page_is_read_separately(self) -> None:
        result = extract_text(
            "application/pdf",
            build_pdf(["Policy number: CP-2026-4471", "Cause of loss: escape of water"]),
        )
        assert result.status is DocumentExtractionStatus.EXTRACTED
        assert result.extractor == "pdf_text_layer"
        assert result.page_count == 2
        assert len(result.pages) == 2
        assert "CP-2026-4471" in result.pages[0].text
        assert "escape of water" in result.pages[1].text

    def test_page_offsets_address_the_joined_text_exactly(self) -> None:
        result = extract_text(
            "application/pdf", build_pdf(["First page body", "Second page body", "Third"])
        )
        for page in result.pages:
            assert result.text[page.char_start : page.char_end] == page.text
        # And the offsets are what gets stored against the document.
        assert result.page_offsets == [[p.char_start, p.char_end] for p in result.pages]

    def test_a_page_with_no_text_layer_is_reported_as_needing_ocr(self) -> None:
        result = extract_text("application/pdf", build_pdf([], blank_pages=2))
        assert result.status is DocumentExtractionStatus.UNSUPPORTED
        assert "optical character recognition" in (result.error or "").lower()
        assert result.page_count == 2

    def test_a_file_that_is_not_a_pdf_fails_without_raising(self) -> None:
        result = extract_text("application/pdf", b"%PDF-1.7\nnowhere near a pdf")
        assert result.status is DocumentExtractionStatus.FAILED
        assert result.error is not None


# --- Word ---------------------------------------------------------------------


class TestWord:
    def test_headings_tables_and_paragraphs_arrive_in_body_order(self) -> None:
        result = extract_text(DOCX_TYPE, build_docx())
        assert result.status is DocumentExtractionStatus.EXTRACTED
        assert result.extractor == "docx"
        assert result.table_count == 2

        lines = result.text.splitlines()
        assert lines[0] == "# Description of loss"
        # The form table's values sit under the first heading, and the matrix under
        # the second — which is only true if the body was walked in order.
        assert lines.index("Policy number: CP-2026-4471") < lines.index("# Schedule")
        assert lines.index("# Schedule") < lines.index("| Item | Qty | Cost |")

    def test_each_heading_starts_a_section_the_officer_can_be_pointed_at(self) -> None:
        result = extract_text(DOCX_TYPE, build_docx())
        assert [page.label for page in result.pages] == ["Description of loss", "Schedule"]
        # Word has no pages until it is laid out, so none are claimed.
        assert result.page_count is None

    def test_a_form_table_is_readable_by_the_deterministic_extractor(self) -> None:
        # The point of rendering a two-column table as `Label: value`: the regex
        # reader that runs when no model is configured can read it.
        result = extract_text(DOCX_TYPE, build_docx())
        reading = extract_from_text(result.text)
        assert reading.policy.policy_number.value == "CP-2026-4471"
        assert reading.policy.insured_name.value == "Harborview Logistics Ltd"

    def test_a_document_the_structured_reader_cannot_open_falls_back(self) -> None:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr(
                "word/document.xml",
                '<?xml version="1.0"?>'
                '<w:document xmlns:w="http://schemas.openxmlformats.org/'
                'wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>'
                "Estimated loss: GBP 128,000</w:t></w:r></w:p></w:body></w:document>",
            )
        result = extract_text(DOCX_TYPE, buffer.getvalue())
        assert result.status is DocumentExtractionStatus.EXTRACTED
        assert result.extractor == "docx_fallback"
        assert "Estimated loss: GBP 128,000" in result.text
        # And it says so, rather than quietly returning worse text.
        assert any("without table layout" in warning for warning in result.warnings)


# --- Excel --------------------------------------------------------------------


class TestExcel:
    def test_each_visible_sheet_is_a_page_named_after_itself(self) -> None:
        result = extract_text(XLSX_TYPE, build_xlsx())
        assert result.status is DocumentExtractionStatus.EXTRACTED
        assert result.extractor == "xlsx"
        assert [page.label for page in result.pages] == ["Summary", "Schedule of values"]
        assert result.text.startswith("## Sheet: Summary")

    def test_a_hidden_sheet_is_not_read(self) -> None:
        # Hidden because its author did not want it read, and a claims file is
        # disclosable.
        result = extract_text(XLSX_TYPE, build_xlsx())
        assert "Do not disclose" not in result.text

    def test_formulas_are_never_emitted_as_text(self) -> None:
        result = extract_text(XLSX_TYPE, build_xlsx())
        assert "=SUM" not in result.text

    def test_dates_read_as_dates_and_whole_floats_lose_their_point(self) -> None:
        result = extract_text(XLSX_TYPE, build_xlsx())
        assert "Date of loss: 2026-03-08" in result.text
        # A currency cell holding 14190.0 is not "14190.0" to a reader, and a date
        # cell is not "2026-03-08 00:00:00".
        assert "| Roof panels | 12 | 14190 | Replaced in full |" in result.text
        assert "14190.0" not in result.text
        assert "00:00:00" not in result.text

    def test_an_empty_column_inside_the_used_range_is_stripped(self) -> None:
        result = extract_text(XLSX_TYPE, build_xlsx())
        schedule = result.pages[1].text
        assert "| Item | Qty | Cost | Notes |" in schedule
        assert "|  |" not in schedule

    def test_blocks_separated_by_blank_rows_are_read_as_separate_tables(self) -> None:
        result = extract_text(XLSX_TYPE, build_xlsx())
        summary = result.pages[0].text
        # Two tables on one sheet, not one table with a hole in it — so the second
        # block's label/value pairing survives instead of being read under the
        # first block's headings.
        assert "Policy number: CP-2026-4471" in summary
        assert "Adjuster: M. Odele" in summary


# --- saved email --------------------------------------------------------------


class TestSavedEmail:
    def test_the_envelope_comes_out_as_labelled_lines(self) -> None:
        result = extract_text("message/rfc822", build_eml())
        assert result.status is DocumentExtractionStatus.EXTRACTED
        assert result.extractor == "eml"
        assert "From: Jane Brooks <jane.brooks@brokerage.example>" in result.text
        assert "Subject: FNOL - burst pipe at Unit 7" in result.text
        assert "Date: Tue, 10 Mar 2026 09:14:00 +0000" in result.text

    def test_an_html_body_is_stripped_to_readable_lines(self) -> None:
        result = extract_text("message/rfc822", build_eml(html=True))
        assert "<p>" not in result.text
        # Block tags became newlines, which is what lets the label reader anchor.
        reading = extract_from_text(result.text)
        assert reading.policy.policy_number.value == "CP-2026-4471"

    def test_an_attachment_is_read_and_its_base64_never_reaches_the_text(self) -> None:
        pdf = build_pdf(["Repair estimate total: GBP 14,190"])
        result = extract_text("message/rfc822", build_eml(attachment=pdf))

        assert "Repair estimate total: GBP 14,190" in result.text
        # The regression this replaces: the whole attachment arrived base64-encoded.
        assert "JVBER" not in result.text
        assert len(result.pages) == 2
        assert result.pages[1].label == "estimate.pdf"

    def test_an_attachment_that_cannot_be_read_is_named_rather_than_dropped(self) -> None:
        message = EmailMessage()
        message["From"] = "broker@example.test"
        message["Subject"] = "Site photographs"
        message.set_content("Photographs attached.")
        message.add_attachment(
            b"\x89PNG\r\n\x1a\nnot really", maintype="image", subtype="png", filename="damage.png"
        )
        result = extract_text("message/rfc822", message.as_bytes())
        assert "Attachments not read: damage.png" in result.text

    def test_a_forwarded_message_is_read_one_level_down(self) -> None:
        inner = EmailMessage()
        inner["From"] = "Jane Brooks <jane.brooks@brokerage.example>"
        inner["Subject"] = "Original notification"
        inner.set_content("Policy number: CP-2026-4471\nDate of loss: 08 March 2026")

        outer = EmailMessage()
        outer["From"] = "handler@carrier.example"
        outer["Subject"] = "FW: chain"
        outer.set_content("See the chain below.")
        outer.add_attachment(inner)

        result = extract_text("message/rfc822", outer.as_bytes())

        assert result.status is DocumentExtractionStatus.EXTRACTED
        assert "See the chain below." in result.pages[0].text
        # The forwarded message's content is read, and read as an attachment rather
        # than being mixed into the covering note's own body.
        assert "CP-2026-4471" not in result.pages[0].text
        assert "CP-2026-4471" in result.text
        assert extract_from_text(result.text).policy.policy_number.value == "CP-2026-4471"

    def test_a_forwarded_message_with_no_filename_is_still_read(self) -> None:
        # Outlook does not always name an attached message, and the extension is what
        # picks the reader — so a nameless one has to be given a `.eml` name.
        inner = EmailMessage()
        inner["Subject"] = "Original notification"
        inner.set_content("Policy number: CP-2026-4471")

        outer = EmailMessage()
        outer["Subject"] = "FW: chain"
        outer.set_content("Forwarded.")
        outer.add_attachment(inner)
        del outer.get_payload()[1]["Content-Disposition"]

        result = extract_text("message/rfc822", outer.as_bytes())
        assert "CP-2026-4471" in result.text

    def test_nesting_stops_one_level_down(self) -> None:
        innermost = EmailMessage()
        innermost["Subject"] = "Deepest"
        innermost.set_content("Policy number: XX-9999-0000")

        middle = EmailMessage()
        middle["Subject"] = "Middle"
        middle.set_content("Middle body.")
        middle.add_attachment(innermost)

        outer = EmailMessage()
        outer["Subject"] = "FW: FW: chain"
        outer.set_content("Outer body.")
        outer.add_attachment(middle)

        result = extract_text("message/rfc822", outer.as_bytes())

        assert "Middle body." in result.text
        # Two levels down is named, never opened — which is what stops a chain of
        # forwards from being followed indefinitely.
        assert "XX-9999-0000" not in result.text

    def test_an_outlook_message_is_read_from_its_mapi_streams(self) -> None:
        content = build_msg(
            {
                "0037": "Motor claim - Vauxhall Vivaro",
                "0C1A": "Jane Brooks",
                "0C1F": "jane.brooks@brokerage.example",
                "0E04": "claims@carrier.example",
                "1000": "Policy number: MC-2026-77410\nDate of loss: 14 March 2026",
            }
        )
        result = extract_text(MSG_TYPE, content)

        assert result.status is DocumentExtractionStatus.EXTRACTED
        assert result.extractor == "msg"
        assert "From: Jane Brooks <jane.brooks@brokerage.example>" in result.text
        assert "Subject: Motor claim - Vauxhall Vivaro" in result.text
        assert result.pages[0].label == "Motor claim - Vauxhall Vivaro"

        reading = extract_from_text(result.text)
        assert reading.policy.policy_number.value == "MC-2026-77410"

    def test_the_nul_padding_outlook_writes_never_reaches_the_text(self) -> None:
        # Postgres rejects NUL in a text column outright, so this is correctness.
        result = extract_text(MSG_TYPE, build_msg({"0037": "Subject", "1000": "Body"}))
        assert "\x00" not in result.text

    def test_a_file_that_is_not_an_outlook_message_is_refused_not_failed(self) -> None:
        result = extract_text(MSG_TYPE, b"Subject: this is an eml, misnamed\n\nbody")
        assert result.status is DocumentExtractionStatus.UNSUPPORTED
        assert "not one" in (result.error or "")

    def test_a_truncated_outlook_container_does_not_raise(self) -> None:
        content = build_msg({"0037": "Subject"})[:600]
        result = extract_text(MSG_TYPE, content)
        assert result.status in (
            DocumentExtractionStatus.FAILED,
            DocumentExtractionStatus.UNSUPPORTED,
        )
        assert result.text == ""


# --- the registry ------------------------------------------------------------


class TestRegistry:
    @pytest.mark.parametrize(
        "content_type",
        [
            "application/pdf",
            DOCX_TYPE,
            XLSX_TYPE,
            "message/rfc822",
            MSG_TYPE,
            "text/csv",
            "application/json",
        ],
    )
    def test_every_claims_format_has_a_reader(self, content_type: str) -> None:
        from app.services.documents.text import registered_extractor

        assert registered_extractor(content_type) is not None

    def test_a_reader_that_raises_is_contained(self) -> None:
        from app.services.documents.text import register_extractor, registered_extractor

        original = registered_extractor("text/plain")

        def explode(content: bytes) -> None:
            raise RuntimeError("boom")

        try:
            register_extractor("text/plain", explode)  # type: ignore[arg-type]
            result = extract_text("text/plain", b"anything")
        finally:
            register_extractor("text/plain", original)  # type: ignore[arg-type]

        assert result.status is DocumentExtractionStatus.FAILED
        assert "RuntimeError" in (result.error or "")
